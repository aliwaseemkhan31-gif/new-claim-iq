"""Render a composed report document to PDF and DOCX.

Both renderers take the plain-data document produced by
:mod:`claimiq.reports.domain.claim_report` and add no content of their own:
what the PDF says and what the screen says are the same document.
"""
from __future__ import annotations

import io
from typing import Any, Mapping
from xml.sax.saxutils import escape

TONE_COLOURS = {
    "neutral": "#F1F3F6",
    "info": "#EEF2FF",
    "warning": "#FFF7E6",
    "danger": "#FDECEC",
}


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------


def render_pdf(document: Mapping[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether,
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9, leading=12, alignment=TA_LEFT)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, leading=10, textColor=colors.HexColor("#4B5563"))
    label = ParagraphStyle("label", parent=body, fontName="Helvetica-Bold")
    heading = ParagraphStyle("heading", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4)
    title = ParagraphStyle("title", parent=styles["Title"], fontSize=16, alignment=TA_LEFT, spaceAfter=2)

    def p(value: Any, style=body) -> Paragraph:
        return Paragraph(escape(str(value if value is not None else "")).replace("\n", "<br/>"), style)

    width = A4[0] - 36 * mm
    story: list = [p(document.get("title", ""), title), p(document.get("subtitle", ""), small), Spacer(1, 4)]

    meta = document.get("meta") or []
    if meta:
        table = Table([[p(k, label), p(v)] for k, v in meta], colWidths=[40 * mm, width - 40 * mm])
        table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        story.append(table)
    if document.get("disclaimer"):
        story += [Spacer(1, 4), _note_table(p(document["disclaimer"], small), "neutral", width, Table, TableStyle, colors)]

    grid = TableStyle(
        [
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]
    )

    for section in document.get("sections") or []:
        story.append(p(section.get("heading", ""), heading))
        for block in section.get("blocks") or []:
            kind = block.get("kind")
            if kind == "paragraph":
                story.append(p(block.get("text")))
            elif kind == "facts":
                table = Table(
                    [[p(k, label), p(v)] for k, v in block.get("facts") or []],
                    colWidths=[45 * mm, width - 45 * mm],
                )
                table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
                story.append(table)
            elif kind == "table":
                columns = block.get("columns") or []
                rows = [[p(c, label) for c in columns]] + [
                    [p(cell, small) for cell in row] for row in block.get("rows") or []
                ]
                table = Table(rows, repeatRows=1, colWidths=[width / max(len(columns), 1)] * len(columns))
                table.setStyle(grid)
                story.append(table)
            elif kind == "note":
                story.append(_note_table(p(block.get("text"), small), block.get("tone", "neutral"), width, Table, TableStyle, colors))
            elif kind == "list":
                story.append(
                    ListFlowable(
                        [ListItem(p(item), leftIndent=10) for item in block.get("items") or []],
                        bulletType="bullet",
                        start="•",
                    )
                )
            elif kind == "findings":
                for finding in block.get("findings") or []:
                    parts = [p(finding.get("statement"), label)]
                    if finding.get("effective_statement"):
                        parts.append(p(f"As amended by reviewer: {finding['effective_statement']}"))
                    parts.append(p(f"{finding.get('epistemic_status', '')} · {finding.get('review', '')}", small))
                    if finding.get("review_reason"):
                        parts.append(p(f"Reviewer's reason: {finding['review_reason']}", small))
                    for citation in finding.get("citations") or []:
                        parts.append(p(f"Source: {citation}", small))
                    tone = "warning" if finding.get("review_state") == "unreviewed" else "neutral"
                    story.append(KeepTogether([_note_table(parts, tone, width, Table, TableStyle, colors)]))
            story.append(Spacer(1, 4))

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(18 * mm, 10 * mm, str(document.get("title", ""))[:110])
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm,
        title=str(document.get("title", "")),
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def _note_table(content, tone: str, width: float, Table, TableStyle, colors):
    cells = content if isinstance(content, list) else [content]
    table = Table([[cells]], colWidths=[width])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(TONE_COLOURS.get(tone, TONE_COLOURS["neutral"]))),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------


def render_docx(document: Mapping[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(str(document.get("title", "")), level=0)
    if document.get("subtitle"):
        doc.add_paragraph(str(document["subtitle"]))

    def facts_table(pairs) -> None:
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, value in pairs:
            row = table.add_row().cells
            row[0].text = str(key)
            row[0].paragraphs[0].runs[0].bold = True
            row[1].text = str(value)

    def note(text: str, tone: str) -> None:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(("Note: " if tone != "danger" else "Attention: ") + text)
        run.italic = True
        if tone in ("warning", "danger"):
            run.font.color.rgb = RGBColor(0x9A, 0x34, 0x12) if tone == "warning" else RGBColor(0xB9, 0x1C, 0x1C)

    if document.get("meta"):
        facts_table(document["meta"])
    if document.get("disclaimer"):
        note(document["disclaimer"], "neutral")

    for section in document.get("sections") or []:
        doc.add_heading(str(section.get("heading", "")), level=1)
        for block in section.get("blocks") or []:
            kind = block.get("kind")
            if kind == "paragraph":
                doc.add_paragraph(str(block.get("text", "")))
            elif kind == "facts":
                facts_table(block.get("facts") or [])
            elif kind == "table":
                columns = block.get("columns") or []
                table = doc.add_table(rows=1, cols=len(columns))
                table.style = "Table Grid"
                for cell, column in zip(table.rows[0].cells, columns):
                    cell.text = str(column)
                    cell.paragraphs[0].runs[0].bold = True
                for values in block.get("rows") or []:
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = str(value)
            elif kind == "note":
                note(str(block.get("text", "")), block.get("tone", "neutral"))
            elif kind == "list":
                for item in block.get("items") or []:
                    doc.add_paragraph(str(item), style="List Bullet")
            elif kind == "findings":
                for finding in block.get("findings") or []:
                    paragraph = doc.add_paragraph(style="List Bullet")
                    paragraph.add_run(str(finding.get("statement", ""))).bold = True
                    if finding.get("effective_statement"):
                        doc.add_paragraph(f"As amended by reviewer: {finding['effective_statement']}")
                    meta = doc.add_paragraph(f"{finding.get('epistemic_status', '')} · {finding.get('review', '')}")
                    meta.runs[0].italic = True
                    if finding.get("review_reason"):
                        doc.add_paragraph(f"Reviewer's reason: {finding['review_reason']}")
                    for citation in finding.get("citations") or []:
                        doc.add_paragraph(f"Source: {citation}")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
